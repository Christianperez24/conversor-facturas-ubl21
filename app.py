"""
Conversor Internacional Masivo de Facturas XML (UBL 2.1) -> Excel & Dashboard
================================================================================
Aplicación Streamlit para procesamiento general e internacional de facturas
electrónicas, notas crédito y notas débito bajo estándar UBL 2.1.
"""

import base64
import html as _html
import inspect
import io
import zipfile
from datetime import datetime

import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

# =============================================================================
# CONFIGURACIÓN DE PÁGINA
# =============================================================================
st.set_page_config(
    page_title="Conversor Internacional XML → Excel",
    page_icon="🌐",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# Compatibilidad entre versiones de Streamlit (width="stretch" vs use_container_width)
STRETCH = (
    {"width": "stretch"}
    if "width" in inspect.signature(st.button).parameters
    else {"use_container_width": True}
)

# =============================================================================
# ESTILOS CSS
# =============================================================================
st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
:root{--ink:#0F172A;--muted:#64748B;--line:#E2E8F0}
.stApp{background:radial-gradient(900px 420px at 50% -8%,#E3EBFB 0%,transparent 70%),#F6F8FC}
.stApp,.stApp p,.stApp label,.stApp h1,.stApp h2,.stApp h3,.stApp h4,.stApp button,.stApp input{
  font-family:'Inter',-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif}
.block-container{max-width:1180px;padding-top:1.2rem;padding-bottom:3rem}
[data-testid="stHeader"]{background:transparent}
[data-testid="stToolbar"],[data-testid="stDecoration"],#MainMenu,footer{display:none!important}

/* ---------- Hero ---------- */
.hero{position:relative;overflow:hidden;border-radius:22px;padding:2.6rem 2.6rem 2.3rem;color:#fff;margin-bottom:1.1rem;
  background:radial-gradient(520px 260px at 92% 8%,rgba(96,165,250,.45),transparent 60%),
  radial-gradient(420px 260px at 0% 100%,rgba(139,92,246,.40),transparent 60%),
  linear-gradient(135deg,#0B1220 0%,#14295F 55%,#1E40AF 100%);
  box-shadow:0 24px 48px -20px rgba(15,23,42,.55)}
.hero::after{content:"";position:absolute;inset:0;pointer-events:none;
  background-image:linear-gradient(rgba(255,255,255,.05) 1px,transparent 1px),linear-gradient(90deg,rgba(255,255,255,.05) 1px,transparent 1px);
  background-size:36px 36px;-webkit-mask-image:linear-gradient(180deg,rgba(0,0,0,.7),transparent);mask-image:linear-gradient(180deg,rgba(0,0,0,.7),transparent)}
.hero>*{position:relative;z-index:1}
.eyebrow{display:inline-block;padding:.3rem .8rem;border-radius:999px;background:rgba(255,255,255,.12);border:1px solid rgba(255,255,255,.22);
  font-size:.72rem;font-weight:700;letter-spacing:.08em;text-transform:uppercase;margin-bottom:1rem}
.hero h1{font-size:2.3rem;font-weight:800;letter-spacing:-.035em;line-height:1.12;margin:0 0 .7rem;padding:0;color:#fff}
.hero h1 span{background:linear-gradient(90deg,#93C5FD,#C4B5FD);-webkit-background-clip:text;background-clip:text;color:transparent}
.hero p{max-width:660px;font-size:1rem;line-height:1.6;margin:0;color:#CBD5E1}
.badges{display:flex;flex-wrap:wrap;gap:.6rem;margin-top:1.4rem}
.badge{padding:.4rem .85rem;border-radius:10px;background:rgba(255,255,255,.10);border:1px solid rgba(255,255,255,.18);font-size:.8rem;font-weight:600;color:#E2E8F0}

/* ---------- Tarjetas (st.container con key="card_*") ---------- */
[class*="st-key-card_"]{background:#fff;border:1px solid var(--line);border-radius:18px;padding:1.3rem 1.5rem;
  box-shadow:0 1px 2px rgba(15,23,42,.04),0 14px 30px -18px rgba(15,23,42,.22)}
.card-head{margin-bottom:.4rem}
.card-title{font-size:1.02rem;font-weight:700;color:var(--ink);letter-spacing:-.01em}
.card-sub{font-size:.82rem;color:var(--muted);margin-top:.15rem}

/* ---------- Secciones ---------- */
.section{display:flex;align-items:center;gap:.8rem;margin:1.9rem 0 .9rem}
.section-icon{width:38px;height:38px;border-radius:11px;display:grid;place-items:center;font-size:1.1rem;background:#E0E9FF}
.section-title{font-size:1.2rem;font-weight:800;color:var(--ink);letter-spacing:-.02em;line-height:1.2}
.section-sub{font-size:.84rem;color:var(--muted)}

/* ---------- Privacidad / avisos ---------- */
.notice{display:flex;align-items:center;gap:.8rem;background:#fff;border:1px solid var(--line);border-left:4px solid #10B981;
  padding:.85rem 1.2rem;border-radius:14px;margin-bottom:1.1rem;font-size:.88rem;color:#334155;box-shadow:0 2px 6px rgba(15,23,42,.03)}
.notice strong{color:#047857}
.notice.warn{border-left-color:#F59E0B}.notice.warn strong{color:#B45309}

/* ---------- Chips de archivos ---------- */
.chips{display:grid;grid-template-columns:repeat(3,1fr);gap:.8rem;margin:.9rem 0 1rem}
.chip{background:#F6F8FC;border:1px solid var(--line);border-radius:12px;padding:.7rem 1rem;display:flex;justify-content:space-between;align-items:center}
.chip b{font-size:1.25rem;color:#1E40AF}
.chip span{font-size:.74rem;font-weight:700;letter-spacing:.07em;text-transform:uppercase;color:var(--muted)}

/* ---------- KPIs ---------- */
.kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:1rem;margin:.5rem 0 1.1rem}
.kpi{background:#fff;border:1px solid var(--line);border-radius:16px;padding:1.1rem 1.2rem;display:flex;gap:.9rem;align-items:flex-start;
  box-shadow:0 1px 2px rgba(15,23,42,.04),0 14px 30px -18px rgba(15,23,42,.22);transition:transform .2s,box-shadow .2s}
.kpi:hover{transform:translateY(-3px);box-shadow:0 18px 34px -16px rgba(15,23,42,.3)}
.kpi-icon{flex:none;width:42px;height:42px;border-radius:12px;display:grid;place-items:center;font-size:1.2rem}
.kpi-blue .kpi-icon{background:#DBEAFE}.kpi-green .kpi-icon{background:#D1FAE5}
.kpi-amber .kpi-icon{background:#FEF3C7}.kpi-violet .kpi-icon{background:#EDE9FE}
.kpi-label{font-size:.7rem;font-weight:700;letter-spacing:.08em;text-transform:uppercase;color:var(--muted)}
.kpi-value{font-size:1.45rem;font-weight:800;color:var(--ink);letter-spacing:-.02em;line-height:1.25;word-break:break-word}
.kpi-value small{font-size:.75rem;font-weight:700;color:var(--muted);margin-right:.25rem}
.kpi-sub{font-size:.76rem;color:var(--muted);margin-top:.15rem}

/* ---------- Uploader ---------- */
[data-testid="stFileUploaderDropzone"]{background:linear-gradient(180deg,#F8FAFF,#F0F5FF);border:2px dashed #BFD0F5!important;
  border-radius:16px!important;padding:1.8rem!important;transition:all .2s}
[data-testid="stFileUploaderDropzone"]:hover{border-color:#2563EB!important;background:#EAF1FF}
[data-testid="stFileUploaderDropzone"] button{border-radius:10px;border:1px solid #BFD0F5;color:#1E40AF;font-weight:600;background:#fff}

/* ---------- Botones ---------- */
button[kind="primary"],button[data-testid="stBaseButton-primary"]{background:linear-gradient(135deg,#2563EB,#4F46E5)!important;
  border:none!important;border-radius:12px!important;padding:.7rem 1.5rem!important;font-weight:700!important;
  box-shadow:0 10px 22px -10px rgba(37,99,235,.8)!important;transition:all .2s}
button[kind="primary"] p,button[data-testid="stBaseButton-primary"] p{color:#fff!important;font-weight:700}
button[kind="primary"]:hover,button[data-testid="stBaseButton-primary"]:hover{transform:translateY(-1px);box-shadow:0 14px 26px -10px rgba(37,99,235,.9)!important}
button[kind="primary"]:disabled,button[data-testid="stBaseButton-primary"]:disabled{opacity:.45;box-shadow:none!important;transform:none}
button[kind="secondary"],button[data-testid="stBaseButton-secondary"]{border-radius:12px!important;border:1px solid var(--line)!important;font-weight:600}

/* ---------- Tabs, tablas, alertas ---------- */
.stTabs [data-baseweb="tab-list"]{gap:.35rem;background:#EAEFF8;padding:.35rem;border-radius:13px;border:none}
.stTabs [data-baseweb="tab"]{height:2.4rem;border-radius:10px;padding:0 1.1rem;font-weight:600;color:var(--muted)}
.stTabs [aria-selected="true"]{background:#fff;color:#1E40AF;box-shadow:0 1px 4px rgba(15,23,42,.12)}
.stTabs [data-baseweb="tab-highlight"],.stTabs [data-baseweb="tab-border"]{display:none}
[data-testid="stDataFrame"]{border:1px solid var(--line);border-radius:12px;overflow:hidden}
[data-testid="stAlert"],[data-testid="stExpander"]{border-radius:14px}

/* ---------- Pasos y pie ---------- */
.steps{display:grid;grid-template-columns:repeat(3,1fr);gap:1rem}
.step{background:#fff;border:1px solid var(--line);border-radius:18px;padding:1.4rem;box-shadow:0 14px 30px -20px rgba(15,23,42,.25)}
.step-n{width:34px;height:34px;border-radius:50%;display:grid;place-items:center;font-weight:800;color:#fff;margin-bottom:.8rem;
  background:linear-gradient(135deg,#2563EB,#7C3AED)}
.step h3{font-size:1.02rem;font-weight:700;color:var(--ink);margin:0 0 .35rem;padding:0}
.step p{font-size:.88rem;line-height:1.55;color:#475569;margin:0}
.footer{text-align:center;color:var(--muted);font-size:.82rem;margin-top:3rem;padding-top:1.4rem;border-top:1px solid var(--line)}
@media(max-width:820px){.kpis,.steps,.chips{grid-template-columns:1fr}.hero{padding:1.8rem 1.4rem}.hero h1{font-size:1.7rem}}
</style>
""",
    unsafe_allow_html=True,
)

# =============================================================================
# HELPERS DE INTERFAZ
# =============================================================================
FONT = "Inter, -apple-system, Segoe UI, sans-serif"
MONEY_COLS = ["Base Gravable", "Impuestos Directos", "Retenciones", "Monto Total",
              "Precio Unitario", "Total Línea", "Cantidad"]


def ui(markup):
    """Renderiza HTML en una sola línea (evita que Markdown lo interprete como código)."""
    st.markdown(" ".join(l.strip() for l in markup.strip().splitlines()), unsafe_allow_html=True)


def section(icon, title, sub=""):
    ui(f'<div class="section"><div class="section-icon">{icon}</div><div>'
       f'<div class="section-title">{title}</div><div class="section-sub">{sub}</div></div></div>')


def card_head(title, sub=""):
    ui(f'<div class="card-head"><div class="card-title">{title}</div><div class="card-sub">{sub}</div></div>')


def card(name):
    """Contenedor tipo tarjeta (usa key para estilizarlo; fallback para versiones antiguas)."""
    try:
        return st.container(key=f"card_{name}")
    except TypeError:
        return st.container(border=True)


def kpi(icon, label, value, sub, tone):
    return (f'<div class="kpi kpi-{tone}"><div class="kpi-icon">{icon}</div><div>'
            f'<div class="kpi-label">{label}</div><div class="kpi-value">{value}</div>'
            f'<div class="kpi-sub">{sub}</div></div></div>')


def show_df(df, height=380):
    cfg = {c: st.column_config.NumberColumn(format="%.2f") for c in MONEY_COLS if c in df.columns}
    try:
        st.dataframe(df, hide_index=True, height=height, column_config=cfg, width="stretch")
    except Exception:
        st.dataframe(df, hide_index=True, height=height, column_config=cfg, use_container_width=True)


# =============================================================================
# GRÁFICOS
# =============================================================================
def _style(fig, height):
    fig.update_layout(
        height=height, margin=dict(l=8, r=8, t=8, b=8),
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
        font=dict(family=FONT, color="#334155", size=12),
        hoverlabel=dict(bgcolor="#0F172A", bordercolor="#0F172A", font=dict(color="#fff", family=FONT, size=12)),
    )
    return fig


def chart_top_proveedores(df, cur):
    d = df.copy()
    d["_p"] = d["Emisor"].fillna(d["ID Fiscal Emisor"]).fillna("Sin identificar")
    g = (d.groupby("_p", as_index=False)["Monto Total"].sum()
         .sort_values("Monto Total", ascending=False).head(5).reset_index(drop=True))
    palette = ["#1D4ED8", "#2563EB", "#3B82F6", "#60A5FA", "#93C5FD"]
    short = g["_p"].map(lambda s: s if len(s) <= 30 else s[:29] + "…")
    fig = go.Figure(go.Bar(
        x=g["Monto Total"], y=short, orientation="h", customdata=g["_p"],
        marker=dict(color=palette[: len(g)]),
        text=[f"{cur} {v:,.0f}" for v in g["Monto Total"]], textposition="outside", cliponaxis=False,
        textfont=dict(size=12, color="#0F172A", family=FONT),
        hovertemplate="<b>%{customdata}</b><br>" + cur + " %{x:,.2f}<extra></extra>",
    ))
    try:
        fig.update_traces(marker_cornerradius=8)
    except ValueError:
        pass
    fig.update_xaxes(range=[0, max(g["Monto Total"].max(), 1) * 1.35], showgrid=True, gridcolor="#EEF2F7",
                     zeroline=False, tickformat="~s", showline=False)
    fig.update_yaxes(autorange="reversed", showgrid=False, automargin=True, ticks="")
    fig.update_layout(bargap=0.45)
    total = d["Monto Total"].sum()
    share = (g["Monto Total"].iloc[0] / total * 100) if total else 0
    sub = (f"{d['_p'].nunique()} proveedor(es) en el lote · el principal concentra {share:.0f}% del monto")
    return _style(fig, max(230, 70 * len(g) + 90)), sub


def chart_base_vs_impuestos(base, imp, cur):
    fig = go.Figure(go.Pie(
        labels=["Base gravable", "Impuestos"], values=[base, imp], hole=0.68, sort=False, pull=[0, 0.03],
        marker=dict(colors=["#2563EB", "#10B981"], line=dict(color="#FFFFFF", width=4)),
        textinfo="percent", textfont=dict(size=13, color="#FFFFFF", family=FONT),
        hovertemplate="<b>%{label}</b><br>" + cur + " %{value:,.2f}<br>%{percent}<extra></extra>",
    ))
    fig.add_annotation(
        text=f"<span style='font-size:11px;color:#64748B'>TOTAL</span><br>"
             f"<b style='font-size:17px;color:#0F172A'>{cur} {base + imp:,.0f}</b>",
        x=0.5, y=0.5, showarrow=False)
    fig.update_layout(showlegend=True, legend=dict(orientation="h", yanchor="top", y=-0.02, xanchor="center", x=0.5))
    return _style(fig, 320)


# =============================================================================
# NAMESPACES ESTÁNDAR UBL 2.1 (GLOBAL)
# =============================================================================
NS = {
    "inv": "urn:oasis:names:specification:ubl:schema:xsd:Invoice-2",
    "att": "urn:oasis:names:specification:ubl:schema:xsd:AttachedDocument-2",
    "cac": "urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2",
    "cbc": "urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2",
    "ds":  "http://www.w3.org/2000/09/xmldsig#",
}

# =============================================================================
# HELPERS DE PARSING GENERAL
# =============================================================================
def _text(elem, path):
    if elem is None:
        return None
    node = elem.find(path, NS)
    if node is not None and node.text:
        return node.text.strip()
    return None

def _float(val):
    if val is None:
        return 0.0
    try:
        return float(str(val).replace(",", "."))
    except (ValueError, TypeError):
        return 0.0

def _clean_xml_bytes(xml_bytes):
    if xml_bytes is None:
        return b""
    data = xml_bytes
    for bom in (b"\xef\xbb\xbf", b"\xff\xfe", b"\xfe\xff"):
        if data.startswith(bom):
            data = data[len(bom):]
            break
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1", errors="replace")
    lt_idx = text.find("<")
    if lt_idx > 0:
        text = text[lt_idx:]
    return text.encode("utf-8")

def _safe_parse(xml_bytes):
    if not xml_bytes:
        return None, "Archivo vacío."
    try:
        return etree.fromstring(xml_bytes), None
    except Exception as e_lxml:
        lxml_err = str(e_lxml)
    try:
        import xml.etree.ElementTree as ET
        return ET.fromstring(xml_bytes), None
    except Exception as e_et:
        return None, f"XML mal formado: {e_et} (lxml: {lxml_err})"

def extract_invoice_from_attached(xml_bytes):
    cleaned = _clean_xml_bytes(xml_bytes)
    if not cleaned:
        return None, "Archivo vacío o sin contenido XML."
    root, parse_err = _safe_parse(cleaned)
    if root is None:
        return None, parse_err
    try:
        tag = etree.QName(root).localname
    except (ValueError, AttributeError):
        raw_tag = root.tag
        tag = raw_tag.split("}")[-1] if "}" in raw_tag else raw_tag

    if tag != "AttachedDocument":
        return cleaned, None

    embedded = root.find(".//cac:Attachment/cbc:EmbeddedDocumentBinaryObject", NS)
    if embedded is not None and embedded.text:
        try:
            decoded = base64.b64decode(embedded.text)
            return _clean_xml_bytes(decoded), None
        except Exception as e:
            return None, f"AttachedDocument con base64 inválido: {e}"

    emb_doc = root.find(".//cac:Attachment/cbc:EmbeddedDocument", NS)
    if emb_doc is not None and emb_doc.text:
        try:
            decoded = base64.b64decode(emb_doc.text)
            return _clean_xml_bytes(decoded), None
        except Exception as e:
            return None, f"AttachedDocument con base64 inválido: {e}"

    return None, "Contenedor de documento sin archivo adjunto embebido."

def parse_invoice(xml_bytes, filename):
    invoice_bytes, unpack_err = extract_invoice_from_attached(xml_bytes)
    if invoice_bytes is None:
        raise ValueError(unpack_err or "No se pudo interpretar la estructura XML.")

    root, parse_err = _safe_parse(invoice_bytes)
    if root is None:
        raise ValueError(parse_err or "XML mal formado.")

    try:
        tag = etree.QName(root).localname
    except (ValueError, AttributeError):
        raw_tag = root.tag
        tag = raw_tag.split("}")[-1] if "}" in raw_tag else raw_tag

    TIPOS_VALIDOS = {"Invoice", "CreditNote", "DebitNote"}
    if tag not in TIPOS_VALIDOS:
        raise ValueError(f"Formato no soportado: '{tag}'.")

    emisor_nombre = _text(root, "cac:AccountingSupplierParty/cac:Party/cac:PartyLegalEntity/cbc:RegistrationName")
    emisor_id = _text(root, "cac:AccountingSupplierParty/cac:Party/cac:PartyIdentification/cbc:ID")
    recp_nombre = _text(root, "cac:AccountingCustomerParty/cac:Party/cac:PartyLegalEntity/cbc:RegistrationName")
    recp_id = _text(root, "cac:AccountingCustomerParty/cac:Party/cac:PartyIdentification/cbc:ID")
    factura_id = _text(root, "cbc:ID")
    fecha = _text(root, "cbc:IssueDate")
    subtotal = _float(_text(root, "cac:LegalMonetaryTotal/cbc:LineExtensionAmount"))
    payable = _float(_text(root, "cac:LegalMonetaryTotal/cbc:PayableAmount"))

    moneda_elem = root.find("cbc:DocumentCurrencyCode", NS)
    moneda = moneda_elem.text.strip() if moneda_elem is not None and moneda_elem.text else "USD"

    prefijo = ""
    folio = factura_id or ""
    if factura_id and "-" in factura_id:
        parts = factura_id.split("-", 1)
        prefijo = parts[0]
        folio = parts[1]

    impuestos = 0.0
    retenciones = 0.0

    for tax_total in root.findall("cac:TaxTotal", NS):
        tax_amount = _float(_text(tax_total, "cbc:TaxAmount"))
        first_sub = tax_total.find("cac:TaxSubtotal", NS)
        if first_sub is None:
            continue
        scheme = _text(first_sub, "cac:TaxCategory/cac:TaxScheme/cbc:ID") or ""
        name = (_text(first_sub, "cac:TaxCategory/cac:TaxScheme/cbc:Name") or "").lower()

        if "ret" in name or scheme in ("05", "06", "07"):
            retenciones += tax_amount
        else:
            impuestos += tax_amount

    tipo_doc_label = {
        "Invoice": "Factura",
        "CreditNote": "Nota Crédito",
        "DebitNote": "Nota Débito",
    }.get(tag, tag)

    cabecera = {
        "Archivo": filename,
        "Tipo Doc": tipo_doc_label,
        "Serie/Prefijo": prefijo,
        "Folio/Número": folio,
        "ID Documento": factura_id,
        "Fecha Emisión": fecha,
        "Emisor": emisor_nombre,
        "ID Fiscal Emisor": emisor_id,
        "Receptor": recp_nombre,
        "ID Fiscal Receptor": recp_id,
        "Moneda": moneda,
        "Base Gravable": subtotal,
        "Impuestos Directos": impuestos,
        "Retenciones": retenciones,
        "Monto Total": payable,
    }

    line_tags = ["cac:InvoiceLine", "cac:CreditNoteLine", "cac:DebitNoteLine"]
    items = []
    for line_tag in line_tags:
        for line in root.findall(line_tag, NS):
            line_id = _text(line, "cbc:ID")
            descripcion = _text(line, "cac:Item/cbc:Description")
            cantidad = (
                _float(_text(line, "cbc:InvoicedQuantity"))
                or _float(_text(line, "cbc:CreditedQuantity"))
                or _float(_text(line, "cbc:DebitedQuantity"))
            )
            precio = _float(_text(line, "cac:Price/cbc:PriceAmount"))
            total_linea = _float(_text(line, "cbc:LineExtensionAmount"))
            codigo = _text(line, "cac:Item/cac:SellersItemIdentification/cbc:ID")

            items.append({
                "ID Documento": factura_id,
                "Línea": line_id,
                "Código SK": codigo,
                "Descripción Item": descripcion,
                "Cantidad": cantidad,
                "Precio Unitario": precio,
                "Total Línea": total_linea,
            })

    return cabecera, items, []

# =============================================================================
# PROCESAMIENTO DE ARCHIVOS (SOPORTE XML Y ZIP)
# =============================================================================
def process_uploaded_files(uploaded_files):
    xml_files = []
    errores = []

    for uf in uploaded_files:
        name = uf.name
        data = uf.getvalue()

        if name.lower().endswith(".zip"):
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as zf:
                    for zinfo in zf.infolist():
                        if zinfo.is_dir():
                            continue
                        if zinfo.filename.lower().endswith(".xml"):
                            try:
                                inner = zf.read(zinfo.filename)
                                xml_files.append((zinfo.filename, inner))
                            except Exception as e:
                                errores.append(f"{zinfo.filename}: error al leer del ZIP ({e})")
            except zipfile.BadZipFile:
                errores.append(f"{name}: archivo ZIP corrupto o inválido.")
        elif name.lower().endswith(".xml"):
            xml_files.append((name, data))
        else:
            errores.append(f"{name}: formato no soportado (use archivos .xml o paquetes .zip).")

    cabeceras = []
    todos_items = []

    for fname, xml_bytes in xml_files:
        try:
            cab, items, errs = parse_invoice(xml_bytes, fname)
            cabeceras.append(cab)
            todos_items.extend(items)
            errores.extend(errs)
        except Exception as e:
            errores.append(f"{fname}: {e}")

    cols_cab = [
        "Archivo", "Tipo Doc", "Serie/Prefijo", "Folio/Número", "ID Documento", "Fecha Emisión",
        "Emisor", "ID Fiscal Emisor", "Receptor", "ID Fiscal Receptor", "Moneda",
        "Base Gravable", "Impuestos Directos", "Retenciones", "Monto Total",
    ]
    cols_det = [
        "ID Documento", "Línea", "Código SK", "Descripción Item",
        "Cantidad", "Precio Unitario", "Total Línea",
    ]

    df_cab = pd.DataFrame(cabeceras, columns=cols_cab) if cabeceras else pd.DataFrame(columns=cols_cab)
    df_det = pd.DataFrame(todos_items, columns=cols_det) if todos_items else pd.DataFrame(columns=cols_det)

    return df_cab, df_det, errores

# =============================================================================
# EXPORTACIÓN A EXCEL UNIFICADA
# =============================================================================
def to_excel(df_cab, df_det):
    output = io.BytesIO()
    wb = Workbook()

    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(border_style="thin", color="E2E8F0")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    money_cols = {"Base Gravable", "Impuestos Directos", "Retenciones", "Monto Total", "Precio Unitario", "Total Línea"}

    def write_sheet(ws, df, title):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(df.columns))
        t = ws.cell(row=1, column=1, value=title)
        t.font = Font(name="Calibri", bold=True, size=14, color="1E40AF")
        t.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 28

        for c_idx, col_name in enumerate(df.columns, start=1):
            c = ws.cell(row=3, column=c_idx, value=col_name)
            c.font = header_font
            c.fill = header_fill
            c.alignment = header_align
            c.border = cell_border
        ws.row_dimensions[3].height = 32

        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=False), start=4):
            for c_idx, value in enumerate(row, start=1):
                col_name = df.columns[c_idx - 1]
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.border = cell_border
                if isinstance(value, (int, float)) and not pd.isna(value):
                    if col_name in money_cols or col_name == "Cantidad":
                        cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        for c_idx, col_name in enumerate(df.columns, start=1):
            max_len = max([len(str(col_name))] + [len(str(v)) for v in df[col_name].head(50).tolist()])
            ws.column_dimensions[get_column_letter(c_idx)].width = min(max_len + 4, 55)

        ws.freeze_panes = "A4"

    ws1 = wb.active
    ws1.title = "Cabeceras_Facturas"
    write_sheet(ws1, df_cab, "Resumen General de Facturas UBL 2.1")

    ws2 = wb.create_sheet("Detalle_Items")
    write_sheet(ws2, df_det, "Detalle de Conceptos / Ítems")

    wb.save(output)
    output.seek(0)
    return output.getvalue()

# =============================================================================
# INTERFAZ PRINCIPAL
# =============================================================================
if "result" not in st.session_state:
    st.session_state["result"] = None

ui("""
<div class="hero">
    <div class="eyebrow">UBL 2.1 · Facturación electrónica internacional</div>
    <h1>Conversor masivo de facturas <span>XML → Excel</span></h1>
    <p>Transforma facturas, notas crédito y notas débito en reportes organizados e inteligencia de negocio, en segundos.</p>
    <div class="badges">
        <div class="badge">⚡ Procesamiento por lotes</div>
        <div class="badge">🌍 Multimoneda</div>
        <div class="badge">📊 Dashboard interactivo</div>
        <div class="badge">🔒 100% local y privado</div>
    </div>
</div>
""")

ui("""
<div class="notice"><span style="font-size:1.3rem">🔒</span><span><strong>Privacidad y seguridad local.</strong>
Tus documentos financieros se procesan en memoria temporal y no se almacenan en servidores externos.</span></div>
""")

# ---------- Carga de archivos ----------
with card("upload"):
    card_head("📁 Carga de archivos", "Arrastra facturas .xml o paquetes comprimidos .zip")
    uploaded = st.file_uploader(
        "Archivos de factura", type=["xml", "zip"], accept_multiple_files=True, label_visibility="collapsed"
    )
    n_xml = sum(1 for f in uploaded if f.name.lower().endswith(".xml")) if uploaded else "—"
    n_zip = sum(1 for f in uploaded if f.name.lower().endswith(".zip")) if uploaded else "—"
    n_tot = len(uploaded) if uploaded else "—"
    ui(f'<div class="chips"><div class="chip"><span>Archivos XML</span><b>{n_xml}</b></div>'
       f'<div class="chip"><span>Archivos ZIP</span><b>{n_zip}</b></div>'
       f'<div class="chip"><span>Total del lote</span><b>{n_tot}</b></div></div>')
    run = st.button("🚀 Procesar facturas y generar reporte", type="primary", disabled=not uploaded, **STRETCH)

if run and uploaded:
    with st.spinner("Procesando estructura UBL 2.1..."):
        df_cab, df_det, errores = process_uploaded_files(uploaded)
        excel = to_excel(df_cab, df_det) if not df_cab.empty else None
    st.session_state["result"] = {"cab": df_cab, "det": df_det, "err": errores, "excel": excel}

# ---------- Resultados (persisten entre reruns, p. ej. al descargar) ----------
res = st.session_state["result"]
if res:
    df_cab, df_det, errores, excel = res["cab"], res["det"], res["err"], res["excel"]

    if df_cab.empty:
        st.warning("No se extrajeron datos válidos. Revisa el panel de advertencias.")
    else:
        monedas = sorted(df_cab["Moneda"].dropna().unique())
        cur = monedas[0] if len(monedas) == 1 else "MIX"
        n_docs = len(df_cab)
        t_monto = df_cab["Monto Total"].sum()
        t_base = df_cab["Base Gravable"].sum()
        t_imp = df_cab["Impuestos Directos"].sum()
        pct_imp = (t_imp / t_base * 100) if t_base else 0

        section("📈", "Resumen del lote", "Indicadores principales de los documentos procesados")
        if len(monedas) > 1:
            ui(f'<div class="notice warn"><span style="font-size:1.3rem">⚠️</span><span><strong>Varias monedas detectadas '
               f'({", ".join(monedas)}).</strong> Los totales suman montos de monedas distintas; revisa el detalle por documento.</span></div>')
        ui('<div class="kpis">'
           + kpi("🧾", "Documentos", f"{n_docs:,}", f"{df_cab['Emisor'].nunique()} emisor(es) distintos", "blue")
           + kpi("💰", "Monto total", f"<small>{cur}</small>{t_monto:,.2f}", f"Promedio {t_monto / n_docs:,.2f} por documento", "green")
           + kpi("🏛️", "Impuestos", f"<small>{cur}</small>{t_imp:,.2f}", f"{pct_imp:.1f}% sobre la base gravable", "amber")
           + kpi("📦", "Ítems", f"{len(df_det):,}", f"{len(df_det) / n_docs:.1f} ítems por documento", "violet")
           + "</div>")

        section("📊", "Análisis y previsualización", "Explora los gráficos y los datos extraídos")
        tab_dash, tab_cab, tab_det = st.tabs(["📈 Dashboard", "📋 Cabeceras", "📦 Detalle de ítems"])

        with tab_dash:
            g1, g2 = st.columns(2, gap="large")
            with g1:
                with card("chart_prov"):
                    fig_bar, sub = chart_top_proveedores(df_cab, cur)
                    card_head("Top 5 proveedores por monto", sub)
                    st.plotly_chart(fig_bar, config={"displayModeBar": False})
            with g2:
                with card("chart_pie"):
                    card_head("Distribución: base vs. impuestos", "Composición del monto facturado")
                    if t_base + t_imp > 0:
                        st.plotly_chart(chart_base_vs_impuestos(t_base, t_imp, cur), config={"displayModeBar": False})
                    else:
                        st.info("No hay montos para graficar.")

        with tab_cab:
            with card("tabla_cab"):
                card_head("Cabeceras de documentos", f"{n_docs:,} registro(s)")
                show_df(df_cab)

        with tab_det:
            with card("tabla_det"):
                card_head("Detalle de conceptos / ítems", f"{len(df_det):,} línea(s)")
                show_df(df_det)

        section("⬇️", "Descarga de reporte unificado", "Excel con dos hojas: cabeceras y detalle de ítems")
        with card("descarga"):
            d1, d2 = st.columns([3, 1], gap="large", vertical_alignment="center")
            with d1:
                st.download_button(
                    label="📥 Descargar reporte en Excel (.xlsx)",
                    data=excel,
                    file_name=f"reporte_facturas_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary",
                    **STRETCH,
                )
            with d2:
                st.button("🧹 Nuevo análisis", on_click=lambda: st.session_state.update(result=None), **STRETCH)

    if errores:
        with st.expander(f"⚠️ {len(errores)} advertencia(s) durante el procesamiento", expanded=False):
            for err in errores:
                st.warning(err)
    elif not df_cab.empty:
        ui('<div class="notice" style="margin-top:1rem"><span style="font-size:1.3rem">✅</span>'
           '<span><strong>Listo.</strong> Todos los archivos se procesaron con éxito.</span></div>')
elif not uploaded:
    st.info("Sube uno o varios archivos .xml o un archivo .ZIP comprimido para iniciar.")

# ---------- Cómo funciona ----------
if not res:
    section("💡", "¿Cómo funciona?", "Tres pasos, sin instalar nada")
    ui("""
    <div class="steps">
        <div class="step"><div class="step-n">1</div><h3>Carga inteligente</h3>
            <p>Sube lotes de facturas en <b>.xml</b> o comprimidas en <b>.zip</b> emitidas bajo el estándar UBL 2.1.</p></div>
        <div class="step"><div class="step-n">2</div><h3>Análisis automatizado</h3>
            <p>Extrae datos fiscales, montos, impuestos y monedas, y desglosa cada ítem sin guardar tus datos.</p></div>
        <div class="step"><div class="step-n">3</div><h3>Excel y analítica</h3>
            <p>Exporta una hoja de cálculo profesional de dos pestañas y visualiza métricas clave en gráficos interactivos.</p></div>
    </div>
    """)

ui(f'<div class="footer">Conversor Internacional UBL 2.1 · Procesamiento 100% local y privado · © {datetime.now().year}</div>')
